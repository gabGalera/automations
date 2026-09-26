from __future__ import annotations

import time as time_mod
from collections.abc import Callable
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import numbers

from consolidacao.motor import (
    COLUNAS_CONFIRMACAO_MP,
    COLUNAS_RECEBIVEL,
    COLUNAS_TRANSACAO,
    COL_CONFIRMACAO,
    ID_ADQUIRENTE,
)

ABA_CONSOLIDADO = "consolidado"
CABECALHO = COLUNAS_TRANSACAO + COLUNAS_RECEBIVEL + COLUNAS_CONFIRMACAO_MP
COLUNAS_DATA = ("data", "Data Repasse", "Data Recibo MP")
COLUNAS_TEXTO = {
    ID_ADQUIRENTE,
    "ID Transacao",
    "Aut",
    "Cartao",
    "Parcelas",
    "Parcela Recebivel",
    "Total Parcelas",
    "Valor Transacao",
    "Taxa %",
    "Taxa Valor",
    "Valor Liquido",
    "Total Reembolsado",
    "taxa % cliente",
    "taxa valor cliente",
    "Valor Repasse",
    "Cliente",
    "Adquirente",
    "Status",
    "Tipo",
    "Bandeira",
}
RETRY_SEGUNDOS = 60
RETRY_PASSO = 1


class PlanilhaBloqueada(OSError):
    pass


def exportar(
    caminho: Path,
    consolidado: list[dict],
    *,
    agora: Callable[[], float] | None = None,
    dormir: Callable[[float], None] | None = None,
) -> None:
    agora = agora or time_mod.monotonic
    dormir = dormir or time_mod.sleep
    wb = Workbook()
    aba = wb.active
    aba.title = ABA_CONSOLIDADO
    _escrever_consolidado(aba, consolidado)
    _salvar_com_retry(wb, caminho, agora, dormir)


def _eh_lock(exc: BaseException) -> bool:
    if isinstance(exc, PermissionError):
        return True
    if isinstance(exc, OSError) and getattr(exc, "winerror", None) in (5, 32):
        return True
    return False


def _salvar_com_retry(wb: Workbook, caminho: Path, agora, dormir) -> None:
    limite = agora() + RETRY_SEGUNDOS
    while True:
        try:
            wb.save(caminho)
            return
        except OSError as extra:
            if not _eh_lock(extra):
                raise
            if agora() >= limite:
                raise PlanilhaBloqueada(
                    f"Não foi possível gravar {caminho}: {extra}"
                ) from extra
            dormir(RETRY_PASSO)


def _escrever_consolidado(ws, consolidado: list[dict]) -> None:
    for col, nome in enumerate(CABECALHO, start=1):
        ws.cell(1, col, nome)
    for r_i, row in enumerate(consolidado, start=2):
        for c_i, nome in enumerate(CABECALHO, start=1):
            _set_cell(ws.cell(r_i, c_i), nome, row.get(nome, ""))


def _set_cell(cell: Cell, nome: str, valor) -> None:
    if valor == "" or valor is None:
        cell.value = None
        return
    if nome in COLUNAS_DATA and isinstance(valor, date) and not isinstance(valor, datetime):
        cell.value = valor
        cell.number_format = "DD/MM/YYYY"
        return
    if nome == "hora" and isinstance(valor, time):
        cell.value = valor
        cell.number_format = "HH:MM:SS"
        return
    if nome == COL_CONFIRMACAO and isinstance(valor, (int, float)) and not isinstance(valor, bool):
        cell.value = valor
        return
    if nome in COLUNAS_TEXTO:
        cell.value = str(valor)
        cell.number_format = numbers.FORMAT_TEXT
        return
    cell.value = valor
