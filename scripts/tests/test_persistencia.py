from datetime import date, time
from pathlib import Path

from openpyxl import load_workbook

from consolidacao_base_1.motor import apply_lote
from consolidacao_base_1.persistencia import carregar_consolidado, gravar_output
from tests.test_motor import _transacao, _recebivel


def test_xlsx_grava_data_e_hora_nativas(tmp_path: Path):
    caminho = tmp_path / "output.xlsx"
    consolidado = apply_lote([], {"tipo": "transacao", "linhas": [_transacao()]})
    consolidado = apply_lote(
        consolidado,
        {
            "tipo": "recebivel",
            "linhas": [_recebivel(**{"ID Trans. Adquirente": "176784155259"})],
        },
    )
    gravar_output(caminho, consolidado, hostname_permitido="DaniGalera")
    relido, controle = carregar_consolidado(caminho)

    assert relido[0]["data"] == date(2026, 9, 7)
    assert relido[0]["hora"] == time(11, 59, 50)
    assert relido[0]["Data Repasse"] == date(2026, 4, 2)
    assert isinstance(relido[0]["data"], date)
    assert isinstance(relido[0]["hora"], time)
    assert isinstance(relido[0]["Data Repasse"], date)
    assert relido[0]["ID Trans. Adquirente"] == "176784155259"
    assert relido[0]["Valor Transacao"] == "80,00"
    assert relido[0]["Confirmacao MP"] == ""
    assert relido[0]["Data Recibo MP"] == ""
    assert controle["hostname_permitido"] == "DaniGalera"


def test_xlsx_preserva_confirmacao_mp_e_chaves_extras_de_controle(tmp_path: Path):
    caminho = tmp_path / "output.xlsx"
    consolidado = apply_lote([], {"tipo": "transacao", "linhas": [_transacao()]})
    consolidado[0]["Confirmacao MP"] = -25660.54
    consolidado[0]["Data Recibo MP"] = date(2026, 4, 8)
    gravar_output(
        caminho,
        consolidado,
        hostname_permitido="DaniGalera",
        controle={"tuplas_mp": '[["176784155259", -25660.54, "2026-04-08"]]'},
    )
    relido, controle = carregar_consolidado(caminho)

    assert relido[0]["Confirmacao MP"] == -25660.54
    assert isinstance(relido[0]["Confirmacao MP"], float)
    assert relido[0]["Data Recibo MP"] == date(2026, 4, 8)
    assert isinstance(relido[0]["Data Recibo MP"], date)
    assert relido[0]["ID Trans. Adquirente"] == "176784155259"
    assert controle["hostname_permitido"] == "DaniGalera"
    assert controle["tuplas_mp"] == '[["176784155259", -25660.54, "2026-04-08"]]'

    wb = load_workbook(caminho)
    assert list(next(wb["consolidado"].iter_rows(max_row=1, values_only=True)))[-2:] == [
        "Confirmacao MP",
        "Data Recibo MP",
    ]

    consolidado = apply_lote(
        relido,
        {
            "tipo": "transacao",
            "linhas": [_transacao(**{"Cliente": "DEPOIS"})],
        },
    )
    gravar_output(caminho, consolidado, hostname_permitido="DaniGalera", controle=controle)
    relido2, controle2 = carregar_consolidado(caminho)
    assert relido2[0]["Cliente"] == "DEPOIS"
    assert relido2[0]["Confirmacao MP"] == -25660.54
    assert relido2[0]["Data Recibo MP"] == date(2026, 4, 8)
    assert controle2["tuplas_mp"] == controle["tuplas_mp"]
