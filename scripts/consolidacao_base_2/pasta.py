from __future__ import annotations

import os
from pathlib import Path

from openpyxl import load_workbook

ARQUIVO_MP = "Recebimentos_MP.xlsx"


def eh_origem_mp(caminho: Path) -> bool:
    return os.path.normcase(caminho.name) == os.path.normcase(ARQUIVO_MP)


def ler_xlsx_mp(caminho: Path) -> list[dict]:
    wb = load_workbook(caminho, data_only=False)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if all(v is None or v == "" for v in raw):
            continue
        item = {}
        for nome, valor in zip(header, raw):
            if not nome:
                continue
            item[nome] = "" if valor is None else valor
        out.append(item)
    return out
