from __future__ import annotations

import sys
import time as time_mod
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import numbers
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook as WorkbookType

from consolidacao_base_1.motor import (
    COLUNAS_CONFIRMACAO_MP,
    COLUNAS_RECEBIVEL,
    COLUNAS_TRANSACAO,
    ID_ADQUIRENTE,
)

ABA_CONSOLIDADO = "consolidado"
ABA_CONTROLE = "controle"
CABECALHO = COLUNAS_TRANSACAO + COLUNAS_RECEBIVEL + COLUNAS_CONFIRMACAO_MP
COLUNAS_TEXTO = {
    ID_ADQUIRENTE,
    "ID Transacao",
    "Aut",
    "Cartao",
    "Parcelas",
    "Parcela Recebivel",
    "Total Parcelas",
    "Valor Transacao",
    "Taxa %",
    "Taxa Valor",
    "Valor Liquido",
    "Total Reembolsado",
    "taxa % cliente",
    "taxa valor cliente",
    "Valor Repasse",
    "Cliente",
    "Adquirente",
    "Status",
    "Tipo",
    "Bandeira",
}
RETRY_SEGUNDOS = 60
RETRY_PASSO = 1


def gravar_output(
    caminho: Path,
    consolidado: list[dict],
    hostname_permitido: str,
    *,
    controle: dict | None = None,
    agora: callable | None = None,
    dormir: callable | None = None,
) -> None:
    agora = agora or time_mod.monotonic
    dormir = dormir or time_mod.sleep
    wb = Workbook()
    default = wb.active
    default.title = ABA_CONSOLIDADO
    _escrever_consolidado(default, consolidado)
    aba_controle = wb.create_sheet(ABA_CONTROLE)
    _escrever_controle(aba_controle, hostname_permitido, controle)
    _salvar_com_retry(wb, caminho, agora, dormir)


def carregar_consolidado(
    caminho: Path,
    *,
    agora: callable | None = None,
    dormir: callable | None = None,
) -> tuple[list[dict], dict]:
    if not caminho.exists():
        return [], {}
    agora = agora or time_mod.monotonic
    dormir = dormir or time_mod.sleep
    limite = agora() + RETRY_SEGUNDOS
    ultimo: Exception | None = None
    while True:
        try:
            wb = load_workbook(caminho, data_only=True)
            return _ler_consolidado(wb), _ler_controle(wb)
        except InvalidFileException:
            return [], {}
        except OSError as extra:
            if not _eh_lock(extra):
                raise
            ultimo = extra
            if agora() >= limite:
                break
            dormir(RETRY_PASSO)
    print(f"Não foi possível ler {caminho}: {ultimo}", file=sys.stderr, flush=True)
    raise ultimo


def _eh_lock(exc: BaseException) -> bool:
    if isinstance(exc, PermissionError):
        return True
    if isinstance(exc, OSError) and getattr(exc, "winerror", None) in (5, 32):
        return True
    return False


def _salvar_com_retry(wb: WorkbookType, caminho: Path, agora, dormir) -> None:
    limite = agora() + RETRY_SEGUNDOS
    ultimo: Exception | None = None
    while True:
        try:
            wb.save(caminho)
            return
        except OSError as extra:
            if not _eh_lock(extra):
                raise
            ultimo = extra
            if agora() >= limite:
                break
            dormir(RETRY_PASSO)
    print(f"Não foi possível gravar {caminho}: {ultimo}", file=sys.stderr, flush=True)


def _escrever_consolidado(ws, consolidado: list[dict]) -> None:
    for col, nome in enumerate(CABECALHO, start=1):
        ws.cell(1, col, nome)
    for r_i, row in enumerate(consolidado, start=2):
        for c_i, nome in enumerate(CABECALHO, start=1):
            valor = row.get(nome, "")
            cell = ws.cell(r_i, c_i)
            _set_cell(cell, nome, valor)


def _set_cell(cell: Cell, nome: str, valor) -> None:
    if valor == "" or valor is None:
        cell.value = None
        return
    if nome in ("data", "Data Repasse", "Data Recibo MP") and isinstance(valor, date) and not isinstance(valor, datetime):
        cell.value = valor
        cell.number_format = "DD/MM/YYYY"
        return
    if nome == "hora" and isinstance(valor, time):
        cell.value = valor
        cell.number_format = "HH:MM:SS"
        return
    if nome in COLUNAS_TEXTO:
        cell.value = str(valor)
        cell.number_format = numbers.FORMAT_TEXT
        return
    cell.value = valor


def _ler_consolidado(wb: WorkbookType) -> list[dict]:
    if ABA_CONSOLIDADO not in wb.sheetnames:
        return []
    ws = wb[ABA_CONSOLIDADO]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if all(v is None or v == "" for v in raw):
            continue
        item = {}
        for nome, valor in zip(header, raw):
            item[nome] = _de_excel(nome, valor)
        out.append(item)
    return out


def _de_excel(nome: str, valor):
    if valor is None:
        return ""
    if nome in ("data", "Data Repasse", "Data Recibo MP"):
        if isinstance(valor, datetime):
            return valor.date()
        if isinstance(valor, date):
            return valor
    if nome == "hora":
        if isinstance(valor, datetime):
            return valor.time().replace(microsecond=0)
        if isinstance(valor, time):
            return valor.replace(microsecond=0)
    return valor if valor is not None else ""


def _escrever_controle(ws, hostname_permitido: str, controle: dict | None) -> None:
    ws["A1"] = "chave"
    ws["B1"] = "valor"
    dados = dict(controle or {})
    dados["hostname_permitido"] = hostname_permitido
    chaves = ["hostname_permitido"] + [k for k in dados if k != "hostname_permitido"]
    for i, chave in enumerate(chaves, start=2):
        ws.cell(i, 1, chave)
        ws.cell(i, 2, dados[chave])


def _ler_controle(wb: WorkbookType) -> dict:
    if ABA_CONTROLE not in wb.sheetnames:
        return {}
    ws = wb[ABA_CONTROLE]
    dados = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        chave = str(row[0])
        dados[chave] = row[1] if len(row) > 1 else ""
    return dados
