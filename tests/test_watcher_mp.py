from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from consolidacao_base_1.persistencia import carregar_consolidado, gravar_output
from consolidacao_base_2.pasta import eh_origem_mp, ler_xlsx_mp
from consolidacao_base_2.persistencia import vistas_de_controle
from consolidacao_base_2.watcher import _primeira_largada, sincronizar_mp


def _escrever_mp(caminho: Path, linhas: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    colunas = ["ID Trans. Adquirente", "Confirmacao MP", "Data Recibo MP"]
    for c, nome in enumerate(colunas, start=1):
        ws.cell(1, c, nome)
    for r, linha in enumerate(linhas, start=2):
        for c, nome in enumerate(colunas, start=1):
            ws.cell(r, c, linha.get(nome, ""))
    wb.save(caminho)
    wb.close()


def test_so_recebimentos_mp_conta_como_origem():
    assert eh_origem_mp(Path("Recebimentos_MP.xlsx"))
    assert eh_origem_mp(Path("recebimentos_mp.xlsx"))
    assert not eh_origem_mp(Path("output.xlsx"))
    assert not eh_origem_mp(Path("transacoes_x.csv"))
    assert not eh_origem_mp(Path("~$Recebimentos_MP.xlsx"))


def test_ler_xlsx_mp_devolve_dicts_com_cabecalhos_reais(tmp_path: Path):
    origem = tmp_path / "Recebimentos_MP.xlsx"
    _escrever_mp(
        origem,
        [
            {
                "ID Trans. Adquirente": "176784155259",
                "Confirmacao MP": 80,
                "Data Recibo MP": date(2026, 4, 2),
            }
        ],
    )

    lote = ler_xlsx_mp(origem)

    assert lote[0]["ID Trans. Adquirente"] == "176784155259"
    assert lote[0]["Confirmacao MP"] == 80
    assert lote[0]["Data Recibo MP"] == datetime(2026, 4, 2)


def test_primeira_largada_so_semeia_e_nao_altera_consolidado_nem_origem(tmp_path: Path):
    origem = tmp_path / "Recebimentos_MP.xlsx"
    _escrever_mp(
        origem,
        [
            {
                "ID Trans. Adquirente": "176784155259",
                "Confirmacao MP": 80,
                "Data Recibo MP": date(2026, 4, 2),
            }
        ],
    )
    antes = origem.read_bytes()
    destino = tmp_path / "output.xlsx"
    gravar_output(destino, [], "DaniGalera")

    sincronizar_mp(tmp_path, semear=True)

    consolidado, controle = carregar_consolidado(destino)
    assert consolidado == []
    assert vistas_de_controle(controle) == {("176784155259", 80, date(2026, 4, 2))}
    assert origem.read_bytes() == antes


def test_tupla_nova_depois_da_semente_entra_e_repetida_e_noop(tmp_path: Path):
    origem = tmp_path / "Recebimentos_MP.xlsx"
    _escrever_mp(origem, [])
    gravar_output(tmp_path / "output.xlsx", [], "DaniGalera")
    sincronizar_mp(tmp_path, semear=True)

    _escrever_mp(
        origem,
        [
            {
                "ID Trans. Adquirente": "176784155259",
                "Confirmacao MP": 80,
                "Data Recibo MP": date(2026, 4, 2),
            }
        ],
    )
    sincronizar_mp(tmp_path, semear=False)
    consolidado, controle = carregar_consolidado(tmp_path / "output.xlsx")
    assert consolidado[0]["Confirmacao MP"] == 80
    vistas = vistas_de_controle(controle)

    sincronizar_mp(tmp_path, semear=False)
    consolidado2, controle2 = carregar_consolidado(tmp_path / "output.xlsx")
    assert consolidado2 == consolidado
    assert vistas_de_controle(controle2) == vistas


def test_relargada_aplica_tupla_que_chegou_com_o_processo_morto(tmp_path: Path):
    origem = tmp_path / "Recebimentos_MP.xlsx"
    _escrever_mp(
        origem,
        [
            {
                "ID Trans. Adquirente": "176784155259",
                "Confirmacao MP": 80,
                "Data Recibo MP": date(2026, 4, 2),
            }
        ],
    )
    gravar_output(tmp_path / "output.xlsx", [], "DaniGalera")
    sincronizar_mp(tmp_path, semear=_primeira_largada(tmp_path))

    _escrever_mp(
        origem,
        [
            {
                "ID Trans. Adquirente": "176784155259",
                "Confirmacao MP": 80,
                "Data Recibo MP": date(2026, 4, 2),
            },
            {
                "ID Trans. Adquirente": "176784155259",
                "Confirmacao MP": -80,
                "Data Recibo MP": date(2026, 4, 2),
            },
        ],
    )
    sincronizar_mp(tmp_path, semear=_primeira_largada(tmp_path))

    consolidado, controle = carregar_consolidado(tmp_path / "output.xlsx")
    assert [row["Confirmacao MP"] for row in consolidado] == [-80]
    assert vistas_de_controle(controle) == {
        ("176784155259", 80, date(2026, 4, 2)),
        ("176784155259", -80, date(2026, 4, 2)),
    }

